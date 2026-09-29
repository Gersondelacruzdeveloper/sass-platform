from __future__ import annotations

import hashlib
import io
from typing import Any

from openpyxl import Workbook, load_workbook


REQUIRED_SHEETS = {
    "STANDARDS": [
        "standard_code", "title", "category", "description", "priority", "active",
        "area_scope", "position_scope", "source", "source_priority", "status",
        "applicability", "version", "effective_date", "notes",
    ],
    "PROCEDURES": [
        "procedure_code", "title", "version", "area", "status", "purpose",
        "quick_steps", "source_class", "source_priority", "primary_standard_code",
    ],
    "TEMPLATES": ["template_key", "name", "description", "area_scope", "status"],
    "QUESTIONS": [
        "question_key", "template_key", "standard_code", "question", "score_type",
        "weight", "order", "active", "observer_guidance",
    ],
}

ALLOWED_PRIORITIES = {"low", "medium", "high", "critical"}
ALLOWED_CATEGORIES = {"service", "beverage", "culinary", "luxury", "leadership", "hard_rock"}
ALLOWED_SCORE_TYPES = {"yes_no", "score", "text"}


def _clean(value: Any) -> Any:
    if value is None:
        return ""
    if hasattr(value, "isoformat") and not isinstance(value, (str, bytes)):
        try:
            return value.isoformat()
        except TypeError:
            pass
    return value.strip() if isinstance(value, str) else value


def _as_bool(value: Any, default=True) -> bool:
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    return str(value).strip().lower() in {"true", "yes", "y", "1", "si", "sí", "active"}


def _as_int(value: Any, default=0) -> int:
    if value in (None, ""):
        return default
    return int(value)


def _rows(ws, headers):
    actual = [str(_clean(c.value)) for c in ws[1]]
    missing = [h for h in headers if h not in actual]
    if missing:
        raise ValueError(f"{ws.title}: missing columns: {', '.join(missing)}")
    pos = {name: actual.index(name) for name in headers}
    out = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        data = {h: _clean(row[pos[h]]) if pos[h] < len(row) else "" for h in headers}
        if not any(v not in (None, "") for v in data.values()):
            continue
        data["_row"] = row_no
        out.append(data)
    return out


def parse_training_workbook(file_bytes: bytes) -> dict:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    missing_sheets = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing_sheets:
        raise ValueError(f"Missing required sheets: {', '.join(missing_sheets)}")

    payload = {}
    for sheet, headers in REQUIRED_SHEETS.items():
        payload[sheet.lower()] = _rows(wb[sheet], headers)

    manifest = {}
    if "MANIFEST" in wb.sheetnames:
        ws = wb["MANIFEST"]
        for row in ws.iter_rows(min_row=1, values_only=True):
            if len(row) >= 2 and row[0] not in (None, ""):
                manifest[str(row[0]).strip()] = _clean(row[1])
    payload["manifest"] = manifest
    payload["sha256"] = hashlib.sha256(file_bytes).hexdigest()
    return normalize_payload(payload)


def normalize_payload(payload: dict) -> dict:
    for row in payload["standards"]:
        row["standard_code"] = str(row["standard_code"]).strip()
        row["title"] = str(row["title"]).strip()
        row["category"] = str(row["category"]).strip().lower()
        row["priority"] = str(row["priority"]).strip().lower()
        row["active"] = _as_bool(row["active"], True)
        row["source_priority"] = _as_int(row["source_priority"], 100)
        row["status"] = str(row["status"] or "ACTIVE").strip().upper()

    for row in payload["procedures"]:
        row["procedure_code"] = str(row["procedure_code"]).strip()
        row["title"] = str(row["title"]).strip()
        row["source_priority"] = _as_int(row["source_priority"], 100)
        row["status"] = str(row["status"] or "ACTIVE").strip().upper()

    for row in payload["templates"]:
        row["template_key"] = str(row["template_key"]).strip()
        row["name"] = str(row["name"]).strip()
        row["status"] = str(row["status"] or "ACTIVE").strip().upper()

    for row in payload["questions"]:
        row["question_key"] = str(row["question_key"]).strip()
        row["template_key"] = str(row["template_key"]).strip()
        row["standard_code"] = str(row["standard_code"]).strip()
        row["score_type"] = str(row["score_type"]).strip().lower()
        row["weight"] = _as_int(row["weight"], 1)
        row["order"] = _as_int(row["order"], 0)
        row["active"] = _as_bool(row["active"], True)
    return payload


def validate_payload(payload: dict) -> dict:
    errors, warnings = [], []

    def unique(rows, field, sheet):
        seen = {}
        for r in rows:
            key = str(r.get(field, "")).strip()
            if not key:
                errors.append(f"{sheet} row {r['_row']}: {field} is required.")
                continue
            if key in seen:
                errors.append(f"{sheet}: duplicate {field} '{key}' at rows {seen[key]} and {r['_row']}.")
            seen[key] = r["_row"]
        return set(seen)

    standard_codes = unique(payload["standards"], "standard_code", "STANDARDS")
    procedure_codes = unique(payload["procedures"], "procedure_code", "PROCEDURES")
    template_keys = unique(payload["templates"], "template_key", "TEMPLATES")
    question_keys = unique(payload["questions"], "question_key", "QUESTIONS")

    for r in payload["standards"]:
        if not r["title"]:
            errors.append(f"STANDARDS row {r['_row']}: title is required.")
        if r["category"] not in ALLOWED_CATEGORIES:
            errors.append(f"STANDARDS row {r['_row']}: invalid category '{r['category']}'.")
        if r["priority"] not in ALLOWED_PRIORITIES:
            errors.append(f"STANDARDS row {r['_row']}: invalid priority '{r['priority']}'.")
        if r["source_priority"] < 1:
            errors.append(f"STANDARDS row {r['_row']}: source_priority must be >= 1.")

    for r in payload["procedures"]:
        primary = r.get("primary_standard_code")
        if primary and primary not in standard_codes:
            warnings.append(f"PROCEDURES row {r['_row']}: primary standard '{primary}' is not in this workbook; existing database binding will be checked on apply.")
        if r["source_priority"] < 1:
            errors.append(f"PROCEDURES row {r['_row']}: source_priority must be >= 1.")

    for r in payload["questions"]:
        if r["template_key"] not in template_keys:
            errors.append(f"QUESTIONS row {r['_row']}: template_key '{r['template_key']}' does not exist in TEMPLATES.")
        if r["standard_code"] and r["standard_code"] not in standard_codes:
            warnings.append(f"QUESTIONS row {r['_row']}: standard_code '{r['standard_code']}' is not in this workbook; existing database binding will be checked on apply.")
        if r["score_type"] not in ALLOWED_SCORE_TYPES:
            errors.append(f"QUESTIONS row {r['_row']}: invalid score_type '{r['score_type']}'.")
        if r["weight"] < 1:
            errors.append(f"QUESTIONS row {r['_row']}: weight must be >= 1.")
        if r["order"] < 0:
            errors.append(f"QUESTIONS row {r['_row']}: order must be >= 0.")

    for r in payload["procedures"]:
        if r["procedure_code"] == "PR-AYB-01" and r["status"] == "ACTIVE":
            warnings.append("PR-AYB-01 is ACTIVE in the file. For TGPC it should normally remain non-applicable; review before apply.")

    return {
        "valid": not errors,
        "errors": errors,
        "warnings": warnings,
        "counts": {
            "standards": len(payload["standards"]),
            "procedures": len(payload["procedures"]),
            "templates": len(payload["templates"]),
            "questions": len(payload["questions"]),
        },
        "keys": {
            "standards": len(standard_codes),
            "procedures": len(procedure_codes),
            "templates": len(template_keys),
            "questions": len(question_keys),
        },
    }


def build_blank_template_bytes() -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    readme = wb.create_sheet("README")
    readme.append(["TRAINING IMPORT CENTER — GENERIC HOTEL TEMPLATE"])
    readme.append(["Upload this workbook to Training > Import Center. Preview is required before apply."])
    readme.append(["source_priority", "1 = strongest/current master source; larger numbers are lower priority."])

    for sheet, headers in REQUIRED_SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append(headers)
        ws.freeze_panes = "A2"

    manifest = wb.create_sheet("MANIFEST", 1)
    manifest.append(["dataset_key", "hotel_training_v1"])
    manifest.append(["dataset_version", "1.0"])
    manifest.append(["property_name", ""])

    buff = io.BytesIO()
    wb.save(buff)
    return buff.getvalue()
