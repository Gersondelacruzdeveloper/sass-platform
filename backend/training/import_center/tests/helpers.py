from io import BytesIO
from openpyxl import Workbook


def make_workbook(*, standard_priority=1, standard_title="Welcome ≤3 minutes", include_bad_question=False):
    wb = Workbook()
    wb.remove(wb.active)
    manifest = wb.create_sheet("MANIFEST")
    manifest.append(["dataset_key", "test_dataset"])
    manifest.append(["dataset_version", "1.0"])

    ws = wb.create_sheet("STANDARDS")
    ws.append(["standard_code","title","category","description","priority","active","area_scope","position_scope","source","source_priority","status","applicability","version","effective_date","notes"])
    ws.append(["ALC-01", standard_title, "service", "Guest welcomed in time", "high", True, "Restaurant", "Server", "Current standard", standard_priority, "ACTIVE", "HOTEL", "1", "", ""])

    ws = wb.create_sheet("PROCEDURES")
    ws.append(["procedure_code","title","version","area","status","purpose","quick_steps","source_class","source_priority","primary_standard_code"])
    ws.append(["PR-TEST-01","Service sequence","1","Restaurant","ACTIVE","Guide service","Welcome → serve","Procedure",2,"ALC-01"])

    ws = wb.create_sheet("TEMPLATES")
    ws.append(["template_key","name","description","area_scope","status"])
    ws.append(["TMP-ALC","À la carte evaluation","Quick floor observation","Restaurant","ACTIVE"])

    ws = wb.create_sheet("QUESTIONS")
    ws.append(["question_key","template_key","standard_code","question","score_type","weight","order","active","observer_guidance"])
    ws.append(["Q-ALC-01", "MISSING" if include_bad_question else "TMP-ALC", "ALC-01", "Did the employee welcome the guest in time?", "yes_no", 2, 1, True, "Observe objectively"])

    out = BytesIO(); wb.save(out); out.seek(0)
    return out
