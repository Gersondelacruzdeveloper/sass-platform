import io
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from colmado.catalog_excel_images import (
    CatalogImageValidationError,
    extract_embedded_catalog_images,
    find_image_column,
)


def png_bytes(color="green"):
    output = io.BytesIO()
    Image.new("RGB", (40, 40), color=color).save(output, format="PNG")
    return output.getvalue()


def worksheet_with_image(*, coordinate="K2", content=None):
    workbook = Workbook()
    worksheet = workbook.active
    headers = [
        "referencia_interna",
        "codigo_barra",
        "nombre",
        "marca",
        "presentacion",
        "categoria",
        "unidad",
        "modo_venta",
        "unidades_por_caja",
        "activo",
        "imagen",
    ]
    worksheet.append(headers)
    worksheet.append(["", "7460123456789", "Producto", "Marca"])

    image_content = content or png_bytes()
    worksheet.add_image(ExcelImage(io.BytesIO(image_content)), coordinate)

    saved = io.BytesIO()
    workbook.save(saved)
    workbook.close()
    saved.seek(0)

    loaded_workbook = load_workbook(saved)
    return loaded_workbook, loaded_workbook.active, headers


class CatalogExcelImageTests(unittest.TestCase):
    def test_finds_the_spanish_image_column(self):
        self.assertEqual(
            find_image_column(["nombre", "imagen"]),
            2,
        )
        self.assertEqual(
            find_image_column(["name", "photo", "image"]),
            3,
        )
        self.assertIsNone(find_image_column(["nombre", "marca"]))

    def test_extracts_an_image_from_its_product_row(self):
        workbook, worksheet, headers = worksheet_with_image()
        self.addCleanup(workbook.close)

        result = extract_embedded_catalog_images(worksheet, headers)

        self.assertEqual(set(result), {2})
        self.assertEqual(result[2].row_number, 2)
        self.assertEqual(result[2].extension, ".png")
        self.assertTrue(result[2].content.startswith(b"\x89PNG"))
        self.assertEqual(
            result[2].filename_for("7460123456789"),
            "catalogo-7460123456789.png",
        )
        self.assertEqual(
            result[2].filename_for("../Plátano / grande"),
            "catalogo-Pl-tano-grande.png",
        )

    def test_ignores_images_outside_the_image_column(self):
        workbook, worksheet, headers = worksheet_with_image(
            coordinate="J2",
        )
        self.addCleanup(workbook.close)

        result = extract_embedded_catalog_images(worksheet, headers)

        self.assertEqual(result, {})

    def test_rejects_two_images_for_the_same_product(self):
        workbook = Workbook()
        worksheet = workbook.active
        headers = ["nombre", "imagen"]
        worksheet.append(headers)
        worksheet.append(["Producto", ""])
        worksheet.add_image(ExcelImage(io.BytesIO(png_bytes("red"))), "B2")
        worksheet.add_image(ExcelImage(io.BytesIO(png_bytes("blue"))), "B2")

        with self.assertRaises(CatalogImageValidationError) as context:
            extract_embedded_catalog_images(worksheet, headers)

        self.assertEqual(context.exception.row_number, 2)
        self.assertEqual(
            context.exception.message,
            "Solo puede colocar una imagen por producto.",
        )
        workbook.close()

    def test_returns_no_images_when_template_has_no_image_column(self):
        workbook, worksheet, _headers = worksheet_with_image()
        self.addCleanup(workbook.close)

        result = extract_embedded_catalog_images(
            worksheet,
            ["nombre", "marca"],
        )

        self.assertEqual(result, {})
