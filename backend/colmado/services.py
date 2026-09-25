import csv
import io
import unicodedata
import uuid
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from datetime import timedelta

from django.db import IntegrityError, transaction
from django.db.models import Count, DecimalField, F, Q, Sum
from django.db.models.functions import Coalesce
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from .models import (
    CashRegisterSession,
    CatalogImport,
    CreditTransaction,
    Customer,
    CustomerOrder,
    CustomerOrderItem,
    Expense,
    InventoryItem,
    InventoryCountItem,
    InventoryCountSession,
    InventoryMovement,
    MasterProduct,
    PayrollPayment,
    PurchaseOrder,
    PurchaseOrderItem,
    Sale,
    SaleItem,
    Store,
    Storefront,
    Supplier,
    SupplierProduct,
)


MONEY = Decimal("0.01")


def _money(value):
    return Decimal(value).quantize(MONEY, rounding=ROUND_HALF_UP)


@transaction.atomic
def open_cash_register(*, organisation, opened_by, data):
    """Open the only active cash register allowed for one store."""

    store = (
        Store.objects.select_for_update()
        .filter(
            id=data["store_id"],
            organisation=organisation,
            is_active=True,
        )
        .first()
    )
    if store is None:
        raise ValidationError(
            {"store_id": "La sucursal no pertenece a este negocio o está inactiva."}
        )

    if CashRegisterSession.objects.filter(
        organisation=organisation,
        store=store,
        status=CashRegisterSession.OPEN,
    ).exists():
        raise ValidationError(
            {"store_id": "Esta sucursal ya tiene una caja abierta."}
        )

    opening_amount = _money(data["opening_amount"])
    cash_register = CashRegisterSession.objects.create(
        organisation=organisation,
        store=store,
        opened_by=opened_by,
        status=CashRegisterSession.OPEN,
        opening_amount=opening_amount,
        cash_sales=Decimal("0.00"),
        expected_cash=opening_amount,
        note=data.get("note", "").strip(),
    )
    return CashRegisterSession.objects.select_related(
        "store",
        "opened_by",
        "closed_by",
    ).get(pk=cash_register.pk)


@transaction.atomic
def close_cash_register(*, cash_register, closed_by, data):
    """Close a register once and save its final cash difference."""

    locked_register = (
        CashRegisterSession.objects.select_for_update()
        .get(pk=cash_register.pk)
    )
    if locked_register.status == CashRegisterSession.CLOSED:
        raise ValidationError({"detail": "Esta caja ya fue cerrada."})

    cash_sales = (
        Sale.objects.filter(
            cash_register_session=locked_register,
            payment_method=Sale.CASH,
            status=Sale.COMPLETED,
        ).aggregate(
            total=Coalesce(
                Sum("total"),
                Decimal("0.00"),
                output_field=DecimalField(max_digits=14, decimal_places=2),
            )
        )["total"]
    )
    cash_sales = _money(cash_sales)
    expected_cash = _money(locked_register.opening_amount + cash_sales)
    counted_cash = _money(data["counted_cash"])

    locked_register.status = CashRegisterSession.CLOSED
    locked_register.closed_by = closed_by
    locked_register.cash_sales = cash_sales
    locked_register.expected_cash = expected_cash
    locked_register.counted_cash = counted_cash
    locked_register.difference = _money(counted_cash - expected_cash)
    locked_register.closed_at = timezone.now()
    if "note" in data:
        locked_register.note = data["note"].strip()
    locked_register.save(
        update_fields=(
            "status",
            "closed_by",
            "cash_sales",
            "expected_cash",
            "counted_cash",
            "difference",
            "closed_at",
            "note",
            "updated_at",
        )
    )
    return locked_register


def _resolve_sale_lines(*, organisation, store, lines):
    """Resolve barcode/button inputs and combine repeated scans."""

    inventory_queryset = InventoryItem.objects.filter(
        organisation=organisation,
        store=store,
        is_active=True,
        master_product__is_active=True,
    )
    quantities_by_inventory_id = defaultdict(lambda: Decimal("0.000"))

    for position, line in enumerate(lines, start=1):
        if line.get("inventory_item_id"):
            inventory_item = inventory_queryset.filter(
                id=line["inventory_item_id"],
            ).first()
            identifier = f"producto #{line['inventory_item_id']}"
        else:
            inventory_item = inventory_queryset.filter(
                master_product__barcode=line["barcode"],
            ).first()
            identifier = f"código {line['barcode']}"

        if inventory_item is None:
            raise ValidationError(
                {
                    "items": (
                        f"El artículo {position} ({identifier}) no está disponible "
                        "en esta sucursal."
                    )
                }
            )

        quantities_by_inventory_id[inventory_item.id] += line["quantity"]

    return quantities_by_inventory_id


@transaction.atomic
def create_pos_sale(*, organisation, cashier, data):
    """Create one POS sale and deduct its stock in a single transaction."""

    store = Store.objects.filter(
        id=data["store_id"],
        organisation=organisation,
        is_active=True,
    ).first()
    if store is None:
        raise ValidationError(
            {"store_id": "La sucursal no pertenece a este negocio o está inactiva."}
        )

    cash_register = (
        CashRegisterSession.objects.select_for_update()
        .filter(
            organisation=organisation,
            store=store,
            status=CashRegisterSession.OPEN,
        )
        .first()
    )

    payment_method = data["payment_method"]
    customer = None
    if payment_method == Sale.CREDIT:
        customer = (
            Customer.objects.select_for_update()
            .filter(
                id=data["customer_id"],
                organisation=organisation,
                is_active=True,
            )
            .first()
        )
        if customer is None:
            raise ValidationError(
                {"customer_id": "El cliente no pertenece a este negocio o está inactivo."}
            )

    quantities_by_inventory_id = _resolve_sale_lines(
        organisation=organisation,
        store=store,
        lines=data["items"],
    )
    inventory_items = {
        item.id: item
        for item in (
            InventoryItem.objects.select_for_update()
            .select_related("master_product")
            .filter(
                id__in=sorted(quantities_by_inventory_id),
                organisation=organisation,
                store=store,
                is_active=True,
                master_product__is_active=True,
            )
            .order_by("id")
        )
    }

    if len(inventory_items) != len(quantities_by_inventory_id):
        raise ValidationError(
            {"items": "Uno de los productos ya no está disponible. Escanee nuevamente."}
        )

    sale_item_values = []
    subtotal = Decimal("0.00")

    for inventory_item_id in sorted(quantities_by_inventory_id):
        inventory_item = inventory_items[inventory_item_id]
        quantity = quantities_by_inventory_id[inventory_item_id]

        if inventory_item.quantity < quantity:
            raise ValidationError(
                {
                    "items": (
                        f"No hay suficiente {inventory_item.master_product.name}. "
                        f"Disponible: {inventory_item.quantity}."
                    )
                }
            )

        line_total = _money(inventory_item.sale_price * quantity)
        subtotal += line_total
        sale_item_values.append(
            {
                "inventory_item": inventory_item,
                "product_name": str(inventory_item.master_product),
                "barcode": inventory_item.master_product.barcode or "",
                "unit": inventory_item.master_product.unit,
                "quantity": quantity,
                "unit_price": inventory_item.sale_price,
                "unit_cost": inventory_item.cost_price,
                "line_total": line_total,
            }
        )

    subtotal = _money(subtotal)
    discount = _money(data.get("discount", Decimal("0.00")))
    if discount > subtotal:
        raise ValidationError(
            {"discount": "El descuento no puede superar el subtotal."}
        )

    total = _money(subtotal - discount)
    if payment_method == Sale.CASH:
        amount_received = _money(data["amount_received"])
        if amount_received < total:
            raise ValidationError(
                {
                    "amount_received": (
                        f"Faltan RD${_money(total - amount_received)} para completar la venta."
                    )
                }
            )
        change_due = _money(amount_received - total)
    elif payment_method == Sale.CREDIT:
        if total <= 0:
            raise ValidationError(
                {"discount": "Una venta fiada debe tener un total mayor que cero."}
            )
        new_balance = _money(customer.balance + total)
        if customer.credit_limit > 0 and new_balance > customer.credit_limit:
            available = _money(max(customer.credit_limit - customer.balance, 0))
            raise ValidationError(
                {
                    "customer_id": (
                        f"El cliente solo tiene RD${available} de crédito disponible."
                    )
                }
            )
        amount_received = Decimal("0.00")
        change_due = Decimal("0.00")
    else:
        amount_received = total
        change_due = Decimal("0.00")

    sale = Sale.objects.create(
        organisation=organisation,
        store=store,
        cashier=cashier,
        customer=customer,
        cash_register_session=cash_register,
        payment_method=payment_method,
        subtotal=subtotal,
        discount=discount,
        total=total,
        amount_received=amount_received,
        change_due=change_due,
        credit_due_date=data.get("credit_due_date"),
    )

    SaleItem.objects.bulk_create(
        SaleItem(sale=sale, **values) for values in sale_item_values
    )

    for inventory_item_id, quantity in quantities_by_inventory_id.items():
        inventory_item = inventory_items[inventory_item_id]
        inventory_item.quantity -= quantity
        inventory_item.save(update_fields=("quantity", "updated_at"))

    if payment_method == Sale.CREDIT:
        customer.balance = new_balance
        customer.save(update_fields=("balance", "updated_at"))
        CreditTransaction.objects.create(
            organisation=organisation,
            customer=customer,
            store=store,
            sale=sale,
            created_by=cashier,
            transaction_type=CreditTransaction.CHARGE,
            amount=total,
            balance_after=new_balance,
            note=f"Venta fiada {sale.receipt_number}",
        )

    if payment_method == Sale.CASH and cash_register is not None:
        cash_register.cash_sales = _money(cash_register.cash_sales + total)
        cash_register.expected_cash = _money(
            cash_register.opening_amount + cash_register.cash_sales
        )
        cash_register.save(
            update_fields=("cash_sales", "expected_cash", "updated_at")
        )

    return (
        Sale.objects.select_related(
            "store",
            "cashier",
            "customer",
            "cash_register_session",
        )
        .prefetch_related("items")
        .get(pk=sale.pk)
    )


@transaction.atomic
def void_pos_sale(*, sale, voided_by):
    """Void a completed sale once and return every item to inventory."""

    locked_sale = Sale.objects.select_for_update().get(pk=sale.pk)
    if locked_sale.status == Sale.VOIDED:
        raise ValidationError({"detail": "Esta venta ya fue anulada."})

    cash_register = None
    if locked_sale.payment_method == Sale.CASH and locked_sale.cash_register_session_id:
        cash_register = CashRegisterSession.objects.select_for_update().get(
            pk=locked_sale.cash_register_session_id
        )
        if cash_register.status == CashRegisterSession.CLOSED:
            raise ValidationError(
                {"detail": "No puede anular una venta de una caja ya cerrada."}
            )

    sale_items = list(locked_sale.items.order_by("inventory_item_id"))
    inventory_items = {
        item.id: item
        for item in (
            InventoryItem.objects.select_for_update()
            .filter(id__in=[line.inventory_item_id for line in sale_items])
            .order_by("id")
        )
    }

    for line in sale_items:
        inventory_item = inventory_items[line.inventory_item_id]
        inventory_item.quantity += line.quantity
        inventory_item.save(update_fields=("quantity", "updated_at"))

    if locked_sale.payment_method == Sale.CREDIT:
        customer = Customer.objects.select_for_update().get(
            pk=locked_sale.customer_id,
            organisation=locked_sale.organisation,
        )
        customer.balance = _money(max(customer.balance - locked_sale.total, 0))
        customer.save(update_fields=("balance", "updated_at"))
        CreditTransaction.objects.create(
            organisation=locked_sale.organisation,
            customer=customer,
            store=locked_sale.store,
            created_by=voided_by,
            transaction_type=CreditTransaction.PAYMENT,
            amount=locked_sale.total,
            balance_after=customer.balance,
            note=f"Anulación de venta {locked_sale.receipt_number}",
        )

    locked_sale.status = Sale.VOIDED
    locked_sale.voided_at = timezone.now()
    locked_sale.save(update_fields=("status", "voided_at", "updated_at"))

    if cash_register is not None:
        cash_register.cash_sales = _money(
            max(cash_register.cash_sales - locked_sale.total, Decimal("0.00"))
        )
        cash_register.expected_cash = _money(
            cash_register.opening_amount + cash_register.cash_sales
        )
        cash_register.save(
            update_fields=("cash_sales", "expected_cash", "updated_at")
        )

    return (
        Sale.objects.select_related(
            "store",
            "cashier",
            "customer",
            "cash_register_session",
        )
        .prefetch_related("items")
        .get(pk=locked_sale.pk)
    )


@transaction.atomic
def record_credit_payment(*, organisation, customer, created_by, data):
    """Record a partial or full payment without allowing a negative balance."""

    store = Store.objects.filter(
        id=data["store_id"],
        organisation=organisation,
        is_active=True,
    ).first()
    if store is None:
        raise ValidationError(
            {"store_id": "La sucursal no pertenece a este negocio o está inactiva."}
        )

    locked_customer = (
        Customer.objects.select_for_update()
        .filter(
            id=customer.id,
            organisation=organisation,
            is_active=True,
        )
        .first()
    )
    if locked_customer is None:
        raise ValidationError({"detail": "El cliente no está disponible."})

    amount = _money(data["amount"])
    if amount > locked_customer.balance:
        raise ValidationError(
            {
                "amount": (
                    f"El cliente debe RD${locked_customer.balance}. "
                    "El abono no puede ser mayor que la deuda."
                )
            }
        )

    locked_customer.balance = _money(locked_customer.balance - amount)
    locked_customer.save(update_fields=("balance", "updated_at"))

    transaction_entry = CreditTransaction.objects.create(
        organisation=organisation,
        customer=locked_customer,
        store=store,
        created_by=created_by,
        transaction_type=CreditTransaction.PAYMENT,
        amount=amount,
        balance_after=locked_customer.balance,
        note=data.get("note", "").strip(),
    )

    return transaction_entry, locked_customer


@transaction.atomic
def create_purchase_order(*, organisation, created_by, data):
    """Create an ordered restocking request from supplier price records."""

    store = Store.objects.filter(
        id=data["store_id"],
        organisation=organisation,
        is_active=True,
    ).first()
    if store is None:
        raise ValidationError(
            {"store_id": "La sucursal no pertenece a este negocio o está inactiva."}
        )

    supplier = Supplier.objects.filter(
        id=data["supplier_id"],
        organisation=organisation,
        is_active=True,
    ).first()
    if supplier is None:
        raise ValidationError(
            {"supplier_id": "El suplidor no pertenece a este negocio o está inactivo."}
        )

    inventory_ids = [line["inventory_item_id"] for line in data["items"]]
    inventory_items = {
        item.id: item
        for item in InventoryItem.objects.select_related("master_product").filter(
            id__in=inventory_ids,
            organisation=organisation,
            store=store,
            is_active=True,
            master_product__is_active=True,
        )
    }
    if len(inventory_items) != len(inventory_ids):
        raise ValidationError(
            {"items": "Uno de los productos no pertenece a esta sucursal."}
        )

    supplier_products = {
        item.master_product_id: item
        for item in SupplierProduct.objects.select_related("master_product").filter(
            organisation=organisation,
            supplier=supplier,
            master_product_id__in=[
                item.master_product_id for item in inventory_items.values()
            ],
            is_active=True,
        )
    }

    order_lines = []
    total_cost = Decimal("0.00")
    for line in data["items"]:
        inventory_item = inventory_items[line["inventory_item_id"]]
        supplier_product = supplier_products.get(inventory_item.master_product_id)
        if supplier_product is None:
            raise ValidationError(
                {
                    "items": (
                        f"{inventory_item.master_product.name} no está registrado "
                        "con este suplidor."
                    )
                }
            )

        cases = line["cases"]
        line_total = _money(cases * supplier_product.case_cost)
        total_cost += line_total
        order_lines.append(
            {
                "inventory_item": inventory_item,
                "product_name": str(inventory_item.master_product),
                "cases": cases,
                "units_per_case": supplier_product.units_per_case,
                "case_cost": supplier_product.case_cost,
                "line_total": line_total,
            }
        )

    order = PurchaseOrder.objects.create(
        organisation=organisation,
        store=store,
        supplier=supplier,
        created_by=created_by,
        status=PurchaseOrder.ORDERED,
        total_cost=_money(total_cost),
        note=data.get("note", "").strip(),
        ordered_at=timezone.now(),
    )
    PurchaseOrderItem.objects.bulk_create(
        PurchaseOrderItem(purchase_order=order, **line) for line in order_lines
    )
    return (
        PurchaseOrder.objects.select_related("store", "supplier", "created_by")
        .prefetch_related("items__inventory_item__master_product")
        .get(pk=order.pk)
    )


@transaction.atomic
def receive_purchase_order(*, order, received_by):
    """Receive an order once, update stock and record immutable movements."""

    locked_order = (
        PurchaseOrder.objects.select_for_update()
        .select_related("store", "supplier", "created_by")
        .get(pk=order.pk)
    )
    if locked_order.status == PurchaseOrder.RECEIVED:
        raise ValidationError({"detail": "Esta orden ya fue recibida."})
    if locked_order.status == PurchaseOrder.CANCELLED:
        raise ValidationError({"detail": "Una orden cancelada no puede recibirse."})

    order_items = list(
        locked_order.items.select_related("inventory_item__master_product").order_by(
            "inventory_item_id"
        )
    )
    if not order_items:
        raise ValidationError({"detail": "La orden no contiene productos."})

    inventory_items = {
        item.id: item
        for item in InventoryItem.objects.select_for_update()
        .filter(id__in=[line.inventory_item_id for line in order_items])
        .order_by("id")
    }
    movements = []

    for line in order_items:
        inventory_item = inventory_items[line.inventory_item_id]
        quantity_received = line.ordered_quantity
        old_quantity = inventory_item.quantity
        new_quantity = old_quantity + quantity_received
        purchase_unit_cost = line.case_cost / line.units_per_case

        if new_quantity > 0:
            weighted_cost = (
                (old_quantity * inventory_item.cost_price)
                + (quantity_received * purchase_unit_cost)
            ) / new_quantity
            inventory_item.cost_price = _money(weighted_cost)

        inventory_item.quantity = new_quantity
        inventory_item.save(
            update_fields=("quantity", "cost_price", "updated_at")
        )
        line.received_quantity = quantity_received
        line.save(update_fields=("received_quantity",))
        movements.append(
            InventoryMovement(
                organisation=locked_order.organisation,
                store=locked_order.store,
                inventory_item=inventory_item,
                purchase_order=locked_order,
                created_by=received_by,
                movement_type=InventoryMovement.PURCHASE,
                quantity_change=quantity_received,
                quantity_after=new_quantity,
                unit_cost=_money(purchase_unit_cost),
                note=f"Orden recibida {locked_order.order_number}",
            )
        )

    InventoryMovement.objects.bulk_create(movements)
    locked_order.status = PurchaseOrder.RECEIVED
    locked_order.received_at = timezone.now()
    locked_order.save(update_fields=("status", "received_at", "updated_at"))

    return (
        PurchaseOrder.objects.select_related("store", "supplier", "created_by")
        .prefetch_related("items__inventory_item__master_product")
        .get(pk=locked_order.pk)
    )


@transaction.atomic
def cancel_purchase_order(*, order):
    locked_order = PurchaseOrder.objects.select_for_update().get(pk=order.pk)
    if locked_order.status == PurchaseOrder.RECEIVED:
        raise ValidationError({"detail": "Una orden recibida no puede cancelarse."})
    if locked_order.status == PurchaseOrder.CANCELLED:
        raise ValidationError({"detail": "Esta orden ya está cancelada."})
    locked_order.status = PurchaseOrder.CANCELLED
    locked_order.save(update_fields=("status", "updated_at"))
    return (
        PurchaseOrder.objects.select_related("store", "supplier", "created_by")
        .prefetch_related("items__inventory_item__master_product")
        .get(pk=locked_order.pk)
    )


def _customer_order_queryset():
    return (
        CustomerOrder.objects.select_related(
            "store",
            "storefront",
            "customer",
        )
        .prefetch_related("items")
    )


@transaction.atomic
def create_customer_order(*, storefront, data):
    """Create an idempotent public order without trusting client-side prices."""

    locked_storefront = (
        Storefront.objects.select_for_update()
        .select_related("store", "organisation")
        .filter(pk=storefront.pk, is_active=True, store__is_active=True)
        .first()
    )
    if locked_storefront is None:
        raise ValidationError(
            {"detail": "Este colmado no está disponible para pedidos."}
        )

    existing_order = _customer_order_queryset().filter(
        storefront=locked_storefront,
        idempotency_key=data["idempotency_key"],
    ).first()
    if existing_order is not None:
        return existing_order, False

    if not locked_storefront.is_accepting_orders:
        raise ValidationError(
            {"detail": "Este colmado no está recibiendo pedidos en este momento."}
        )

    inventory_ids = [line["inventory_item_id"] for line in data["items"]]
    inventory_items = {
        item.id: item
        for item in InventoryItem.objects.select_related("master_product").filter(
            id__in=inventory_ids,
            organisation=locked_storefront.organisation,
            store=locked_storefront.store,
            is_active=True,
            master_product__is_active=True,
        )
    }
    if len(inventory_items) != len(inventory_ids):
        raise ValidationError(
            {"items": "Uno de los productos ya no está disponible."}
        )

    order_lines = []
    subtotal = Decimal("0.00")
    for line in data["items"]:
        inventory_item = inventory_items[line["inventory_item_id"]]
        quantity = line["quantity"]
        if inventory_item.quantity < quantity:
            raise ValidationError(
                {
                    "items": (
                        f"No hay suficiente {inventory_item.master_product.name}. "
                        f"Disponible: {inventory_item.quantity}."
                    )
                }
            )

        line_total = _money(inventory_item.sale_price * quantity)
        subtotal += line_total
        order_lines.append(
            {
                "inventory_item": inventory_item,
                "product_name": str(inventory_item.master_product),
                "barcode": inventory_item.master_product.barcode or "",
                "unit": inventory_item.master_product.unit,
                "quantity": quantity,
                "unit_price": inventory_item.sale_price,
                "line_total": line_total,
            }
        )

    subtotal = _money(subtotal)
    if subtotal < locked_storefront.minimum_order:
        raise ValidationError(
            {
                "items": (
                    f"El pedido mínimo es RD${locked_storefront.minimum_order}."
                )
            }
        )

    customer_phone = data["customer_phone"].strip()
    customer = Customer.objects.filter(
        organisation=locked_storefront.organisation,
        phone=customer_phone,
    ).first()
    if customer is None:
        customer = Customer.objects.create(
            organisation=locked_storefront.organisation,
            name=data["customer_name"].strip(),
            phone=customer_phone,
            address=data["delivery_address"].strip(),
        )

    delivery_fee = _money(locked_storefront.delivery_fee)
    order = CustomerOrder.objects.create(
        organisation=locked_storefront.organisation,
        store=locked_storefront.store,
        storefront=locked_storefront,
        customer=customer,
        idempotency_key=data["idempotency_key"],
        customer_name=data["customer_name"].strip(),
        customer_phone=customer_phone,
        delivery_address=data["delivery_address"].strip(),
        latitude=data.get("latitude"),
        longitude=data.get("longitude"),
        delivery_notes=data.get("delivery_notes", "").strip(),
        subtotal=subtotal,
        delivery_fee=delivery_fee,
        total=_money(subtotal + delivery_fee),
    )
    CustomerOrderItem.objects.bulk_create(
        CustomerOrderItem(order=order, **line) for line in order_lines
    )
    return _customer_order_queryset().get(pk=order.pk), True


@transaction.atomic
def change_customer_order_status(*, order, changed_by, new_status):
    """Advance one order safely and commit or restore stock exactly once."""

    locked_order = (
        CustomerOrder.objects.select_for_update()
        .select_related("store", "storefront")
        .get(pk=order.pk)
    )

    if locked_order.status == new_status:
        return _customer_order_queryset().get(pk=locked_order.pk)
    if locked_order.status in (CustomerOrder.DELIVERED, CustomerOrder.CANCELLED):
        raise ValidationError({"status": "Este pedido ya está cerrado."})

    allowed_next = {
        CustomerOrder.NEW: CustomerOrder.PREPARING,
        CustomerOrder.PREPARING: CustomerOrder.READY,
        CustomerOrder.READY: CustomerOrder.ON_THE_WAY,
        CustomerOrder.ON_THE_WAY: CustomerOrder.DELIVERED,
    }
    if new_status != CustomerOrder.CANCELLED:
        expected_status = allowed_next.get(locked_order.status)
        if new_status != expected_status:
            expected_label = dict(CustomerOrder.STATUS_CHOICES).get(
                expected_status,
                "el próximo paso",
            )
            raise ValidationError(
                {"status": f"El próximo estado debe ser: {expected_label}."}
            )

    order_items = list(
        locked_order.items.select_related(
            "inventory_item__master_product"
        ).order_by("inventory_item_id")
    )
    if not order_items:
        raise ValidationError({"detail": "El pedido no contiene productos."})

    if new_status == CustomerOrder.PREPARING and not locked_order.inventory_committed:
        inventory_items = {
            item.id: item
            for item in InventoryItem.objects.select_for_update()
            .filter(id__in=[line.inventory_item_id for line in order_items])
            .order_by("id")
        }
        if len(inventory_items) != len(order_items):
            raise ValidationError(
                {"items": "Uno de los productos ya no está disponible."}
            )

        movements = []
        for line in order_items:
            inventory_item = inventory_items[line.inventory_item_id]
            if (
                inventory_item.organisation_id != locked_order.organisation_id
                or inventory_item.store_id != locked_order.store_id
                or not inventory_item.is_active
                or not inventory_item.master_product.is_active
            ):
                raise ValidationError(
                    {"items": f"{line.product_name} ya no está disponible."}
                )
            if inventory_item.quantity < line.quantity:
                raise ValidationError(
                    {
                        "items": (
                            f"No hay suficiente {line.product_name}. "
                            f"Disponible: {inventory_item.quantity}."
                        )
                    }
                )

            inventory_item.quantity -= line.quantity
            inventory_item.save(update_fields=("quantity", "updated_at"))
            movements.append(
                InventoryMovement(
                    organisation=locked_order.organisation,
                    store=locked_order.store,
                    inventory_item=inventory_item,
                    created_by=changed_by,
                    movement_type=InventoryMovement.CUSTOMER_ORDER,
                    quantity_change=-line.quantity,
                    quantity_after=inventory_item.quantity,
                    unit_cost=inventory_item.cost_price,
                    note=f"Pedido {locked_order.order_number}",
                )
            )

        InventoryMovement.objects.bulk_create(movements)
        locked_order.inventory_committed = True
        locked_order.accepted_at = timezone.now()

    if new_status == CustomerOrder.CANCELLED:
        if locked_order.inventory_committed:
            inventory_items = {
                item.id: item
                for item in InventoryItem.objects.select_for_update()
                .filter(id__in=[line.inventory_item_id for line in order_items])
                .order_by("id")
            }
            movements = []
            for line in order_items:
                inventory_item = inventory_items[line.inventory_item_id]
                inventory_item.quantity += line.quantity
                inventory_item.save(update_fields=("quantity", "updated_at"))
                movements.append(
                    InventoryMovement(
                        organisation=locked_order.organisation,
                        store=locked_order.store,
                        inventory_item=inventory_item,
                        created_by=changed_by,
                        movement_type=InventoryMovement.CUSTOMER_ORDER_CANCELLED,
                        quantity_change=line.quantity,
                        quantity_after=inventory_item.quantity,
                        unit_cost=inventory_item.cost_price,
                        note=f"Cancelación del pedido {locked_order.order_number}",
                    )
                )
            InventoryMovement.objects.bulk_create(movements)
            locked_order.inventory_committed = False
        locked_order.cancelled_at = timezone.now()
    elif new_status == CustomerOrder.READY:
        locked_order.ready_at = timezone.now()
    elif new_status == CustomerOrder.ON_THE_WAY:
        locked_order.dispatched_at = timezone.now()
    elif new_status == CustomerOrder.DELIVERED:
        locked_order.delivered_at = timezone.now()

    locked_order.status = new_status
    locked_order.save(
        update_fields=(
            "status",
            "inventory_committed",
            "accepted_at",
            "ready_at",
            "dispatched_at",
            "delivered_at",
            "cancelled_at",
            "updated_at",
        )
    )
    return _customer_order_queryset().get(pk=locked_order.pk)


def _inventory_count_queryset():
    return (
        InventoryCountSession.objects.select_related("store", "created_by")
        .prefetch_related(
            "items__inventory_item__master_product",
            "items__counted_by",
        )
    )


@transaction.atomic
def create_inventory_count(*, organisation, created_by, data):
    """Open one physical count without allowing two open counts per store."""

    store = (
        Store.objects.select_for_update()
        .filter(
            id=data["store_id"],
            organisation=organisation,
            is_active=True,
        )
        .first()
    )
    if store is None:
        raise ValidationError(
            {"store_id": "La sucursal no pertenece a este negocio o está inactiva."}
        )
    if InventoryCountSession.objects.filter(
        organisation=organisation,
        store=store,
        status=InventoryCountSession.OPEN,
    ).exists():
        raise ValidationError(
            {"store_id": "Esta sucursal ya tiene un conteo abierto."}
        )

    inventory_count = InventoryCountSession.objects.create(
        organisation=organisation,
        store=store,
        created_by=created_by,
        note=data.get("note", "").strip(),
    )
    return _inventory_count_queryset().get(pk=inventory_count.pk)


@transaction.atomic
def count_inventory_item(*, inventory_count, counted_by, data):
    """Record the physical total for one scanned or manually selected product."""

    locked_count = (
        InventoryCountSession.objects.select_for_update()
        .select_related("store")
        .get(pk=inventory_count.pk)
    )
    if locked_count.status != InventoryCountSession.OPEN:
        raise ValidationError({"detail": "Este conteo ya está cerrado."})

    inventory_queryset = (
        InventoryItem.objects.select_for_update()
        .select_related("master_product")
        .filter(
            organisation=locked_count.organisation,
            store=locked_count.store,
            is_active=True,
            master_product__is_active=True,
        )
    )
    if data.get("inventory_item_id"):
        inventory_item = inventory_queryset.filter(
            id=data["inventory_item_id"]
        ).first()
    else:
        inventory_item = inventory_queryset.filter(
            master_product__barcode=data["barcode"]
        ).first()

    if inventory_item is None:
        raise ValidationError(
            {"product": "Este producto no está disponible en esta sucursal."}
        )

    count_item, created = InventoryCountItem.objects.get_or_create(
        inventory_count=locked_count,
        inventory_item=inventory_item,
        defaults={
            "counted_by": counted_by,
            "expected_quantity": inventory_item.quantity,
            "counted_quantity": data["counted_quantity"],
            "reason": data["reason"],
            "note": data.get("note", "").strip(),
        },
    )
    if not created:
        count_item.counted_by = counted_by
        count_item.counted_quantity = data["counted_quantity"]
        count_item.reason = data["reason"]
        count_item.note = data.get("note", "").strip()
        count_item.save(
            update_fields=(
                "counted_by",
                "counted_quantity",
                "reason",
                "note",
                "counted_at",
            )
        )
    return (
        InventoryCountItem.objects.select_related(
            "inventory_item__master_product",
            "counted_by",
        ).get(pk=count_item.pk)
    )


@transaction.atomic
def complete_inventory_count(*, inventory_count, completed_by):
    """Apply every counted total and record immutable inventory movements."""

    locked_count = (
        InventoryCountSession.objects.select_for_update()
        .select_related("store", "created_by")
        .get(pk=inventory_count.pk)
    )
    if locked_count.status == InventoryCountSession.COMPLETED:
        raise ValidationError({"detail": "Este conteo ya fue completado."})
    if locked_count.status == InventoryCountSession.CANCELLED:
        raise ValidationError({"detail": "Un conteo cancelado no puede completarse."})

    count_items = list(
        locked_count.items.select_related(
            "inventory_item__master_product",
            "counted_by",
        ).order_by("inventory_item_id")
    )
    if not count_items:
        raise ValidationError(
            {"detail": "Debe contar al menos un producto antes de completar."}
        )

    inventory_items = {
        item.id: item
        for item in InventoryItem.objects.select_for_update()
        .filter(id__in=[line.inventory_item_id for line in count_items])
        .order_by("id")
    }
    movement_type_by_reason = {
        InventoryCountItem.REGULAR_COUNT: InventoryMovement.INVENTORY_COUNT,
        InventoryCountItem.DAMAGED: InventoryMovement.DAMAGED,
        InventoryCountItem.EXPIRED: InventoryMovement.EXPIRED,
        InventoryCountItem.LOSS: InventoryMovement.LOSS,
        InventoryCountItem.PERSONAL_USE: InventoryMovement.PERSONAL_USE,
        InventoryCountItem.OTHER: InventoryMovement.CORRECTION,
    }
    movements = []
    for count_item in count_items:
        inventory_item = inventory_items[count_item.inventory_item_id]
        quantity_change = count_item.counted_quantity - inventory_item.quantity
        if quantity_change == 0:
            continue

        inventory_item.quantity = count_item.counted_quantity
        inventory_item.save(update_fields=("quantity", "updated_at"))
        reason_label = count_item.get_reason_display()
        detail = count_item.note or reason_label
        movements.append(
            InventoryMovement(
                organisation=locked_count.organisation,
                store=locked_count.store,
                inventory_item=inventory_item,
                inventory_count=locked_count,
                created_by=completed_by,
                movement_type=movement_type_by_reason[count_item.reason],
                quantity_change=quantity_change,
                quantity_after=count_item.counted_quantity,
                unit_cost=inventory_item.cost_price,
                note=f"Conteo {locked_count.count_number}: {detail}",
            )
        )

    InventoryMovement.objects.bulk_create(movements)
    locked_count.status = InventoryCountSession.COMPLETED
    locked_count.completed_at = timezone.now()
    locked_count.save(update_fields=("status", "completed_at", "updated_at"))
    return _inventory_count_queryset().get(pk=locked_count.pk)


@transaction.atomic
def cancel_inventory_count(*, inventory_count):
    """Close an unfinished count without changing any stock."""

    locked_count = InventoryCountSession.objects.select_for_update().get(
        pk=inventory_count.pk
    )
    if locked_count.status == InventoryCountSession.COMPLETED:
        raise ValidationError({"detail": "Un conteo completado no puede cancelarse."})
    if locked_count.status == InventoryCountSession.CANCELLED:
        raise ValidationError({"detail": "Este conteo ya fue cancelado."})

    locked_count.status = InventoryCountSession.CANCELLED
    locked_count.cancelled_at = timezone.now()
    locked_count.save(update_fields=("status", "cancelled_at", "updated_at"))
    return _inventory_count_queryset().get(pk=locked_count.pk)


@transaction.atomic
def adjust_inventory(*, organisation, inventory_item, adjusted_by, data):
    """Set a product's real quantity and leave an auditable reason."""

    locked_item = (
        InventoryItem.objects.select_for_update()
        .select_related("store", "master_product")
        .filter(
            pk=inventory_item.pk,
            organisation=organisation,
            is_active=True,
            store__is_active=True,
            master_product__is_active=True,
        )
        .first()
    )
    if locked_item is None:
        raise ValidationError({"detail": "Este producto no está disponible."})

    new_quantity = data["new_quantity"]
    quantity_change = new_quantity - locked_item.quantity
    if quantity_change == 0:
        raise ValidationError(
            {"new_quantity": "La existencia ya tiene esa cantidad."}
        )
    reduction_reasons = {
        InventoryMovement.DAMAGED,
        InventoryMovement.EXPIRED,
        InventoryMovement.LOSS,
        InventoryMovement.PERSONAL_USE,
    }
    if data["reason"] in reduction_reasons and quantity_change > 0:
        raise ValidationError(
            {"new_quantity": "Este motivo solamente puede reducir la existencia."}
        )

    previous_quantity = locked_item.quantity
    locked_item.quantity = new_quantity
    locked_item.save(update_fields=("quantity", "updated_at"))
    note = data.get("note", "").strip()
    movement = InventoryMovement.objects.create(
        organisation=organisation,
        store=locked_item.store,
        inventory_item=locked_item,
        created_by=adjusted_by,
        movement_type=data["reason"],
        quantity_change=quantity_change,
        quantity_after=new_quantity,
        unit_cost=locked_item.cost_price,
        note=note or f"Ajuste desde {previous_quantity} hasta {new_quantity}",
    )
    return locked_item, movement


CATALOG_MAX_ROWS = 20000
CATALOG_MAX_REPORTED_ERRORS = 200
CATALOG_HEADERS = {
    "internal_reference": (
        "referencia_interna",
        "internal_reference",
        "referencia",
    ),
    "barcode": ("codigo_barra", "barcode", "codigo"),
    "name": ("nombre", "name", "producto"),
    "brand": ("marca", "brand"),
    "presentation": ("presentacion", "presentation"),
    "category": ("categoria", "category"),
    "unit": ("unidad", "unit"),
    "sale_mode": ("modo_venta", "sale_mode"),
    "units_per_case": ("unidades_por_caja", "units_per_case"),
    "is_active": ("activo", "is_active"),
}


class CatalogImportAbort(Exception):
    def __init__(self, errors):
        self.errors = errors
        super().__init__("Catalog import aborted")


def _normalize_header(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    text = "".join(character for character in text if not unicodedata.combining(character))
    return "_".join(text.replace("-", " ").split())


def _clean_cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _canonicalize_catalog_row(raw_row):
    normalized = {
        _normalize_header(key): _clean_cell(value)
        for key, value in raw_row.items()
        if key is not None
    }
    row = {}
    for canonical_name, aliases in CATALOG_HEADERS.items():
        row[canonical_name] = next(
            (normalized[alias] for alias in aliases if alias in normalized),
            "",
        )
    return row


def _read_csv_catalog(file_bytes):
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CatalogImportAbort(
            [{"row": 1, "errors": ["El CSV debe estar guardado en formato UTF-8."]}]
        ) from exc

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if reader.fieldnames is None:
        raise CatalogImportAbort(
            [{"row": 1, "errors": ["El archivo no contiene encabezados."]}]
        )
    return list(reader), reader.fieldnames


def _read_xlsx_catalog(file_bytes):
    try:
        workbook = load_workbook(
            filename=io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if headers is None:
            workbook.close()
            raise CatalogImportAbort(
                [{"row": 1, "errors": ["El archivo de Excel está vacío."]}]
            )
        fieldnames = [_clean_cell(value) for value in headers]
        rows = [dict(zip(fieldnames, values)) for values in iterator]
        workbook.close()
        return rows, fieldnames
    except CatalogImportAbort:
        raise
    except Exception as exc:
        raise CatalogImportAbort(
            [{"row": 1, "errors": ["No se pudo leer el archivo de Excel."]}]
        ) from exc


def _parse_boolean(value, *, default=True):
    normalized = _normalize_header(value)
    if normalized == "":
        return default
    if normalized in {"1", "true", "si", "yes", "activo", "active"}:
        return True
    if normalized in {"0", "false", "no", "inactivo", "inactive"}:
        return False
    raise ValueError("Use Sí/No, True/False o 1/0.")


def _parse_catalog_row(row, row_number):
    errors = []
    name = row["name"].strip()
    if not name:
        errors.append("El nombre es obligatorio.")

    barcode = row["barcode"].strip() or None
    if barcode:
        if not barcode.isdigit():
            errors.append("El código de barra debe contener solamente números.")
        elif len(barcode) not in (8, 12, 13, 14):
            errors.append("El código de barra debe tener 8, 12, 13 o 14 dígitos.")

    internal_reference = None
    if row["internal_reference"]:
        try:
            internal_reference = uuid.UUID(row["internal_reference"])
        except (ValueError, AttributeError):
            errors.append("La referencia interna no es un UUID válido.")

    unit_aliases = {
        "": MasterProduct.UNIT,
        "unit": MasterProduct.UNIT,
        "unidad": MasterProduct.UNIT,
        "lb": MasterProduct.POUND,
        "libra": MasterProduct.POUND,
        "kg": MasterProduct.KILOGRAM,
        "kilogramo": MasterProduct.KILOGRAM,
        "liter": MasterProduct.LITER,
        "litro": MasterProduct.LITER,
        "portion": MasterProduct.PORTION,
        "porcion": MasterProduct.PORTION,
    }
    unit_key = _normalize_header(row["unit"])
    unit = unit_aliases.get(unit_key)
    if unit is None:
        errors.append("La unidad no es válida.")

    sale_mode_aliases = {
        "": MasterProduct.BY_UNIT,
        "unit": MasterProduct.BY_UNIT,
        "unidad": MasterProduct.BY_UNIT,
        "por_unidad": MasterProduct.BY_UNIT,
        "weight": MasterProduct.BY_WEIGHT,
        "peso": MasterProduct.BY_WEIGHT,
        "por_peso": MasterProduct.BY_WEIGHT,
        "amount": MasterProduct.BY_AMOUNT,
        "monto": MasterProduct.BY_AMOUNT,
        "por_monto": MasterProduct.BY_AMOUNT,
        "valor": MasterProduct.BY_AMOUNT,
    }
    sale_mode_key = _normalize_header(row["sale_mode"])
    sale_mode = sale_mode_aliases.get(sale_mode_key)
    if sale_mode is None:
        errors.append("El modo de venta no es válido.")

    try:
        units_per_case = int(row["units_per_case"] or "1")
        if units_per_case < 1:
            raise ValueError
    except (TypeError, ValueError):
        units_per_case = None
        errors.append("Las unidades por caja deben ser un número entero mayor que cero.")

    try:
        is_active = _parse_boolean(row["is_active"])
    except ValueError as exc:
        is_active = None
        errors.append(str(exc))

    if errors:
        return None, {"row": row_number, "errors": errors}

    return {
        "row_number": row_number,
        "internal_reference": internal_reference,
        "barcode": barcode,
        "name": name,
        "brand": row["brand"].strip(),
        "presentation": row["presentation"].strip(),
        "category": row["category"].strip(),
        "unit": unit,
        "sale_mode": sale_mode,
        "units_per_case": units_per_case,
        "is_active": is_active,
    }, None


def _validate_catalog_rows(raw_rows, fieldnames):
    normalized_headers = {_normalize_header(value) for value in fieldnames}
    name_aliases = set(CATALOG_HEADERS["name"])
    if not normalized_headers.intersection(name_aliases):
        return [], [
            {
                "row": 1,
                "errors": ["Falta la columna obligatoria: nombre."],
            }
        ]

    nonempty_rows = []
    for index, raw_row in enumerate(raw_rows, start=2):
        canonical = _canonicalize_catalog_row(raw_row)
        if any(canonical.values()):
            nonempty_rows.append((index, canonical))
    if len(nonempty_rows) > CATALOG_MAX_ROWS:
        return [], [
            {
                "row": 1,
                "errors": [f"El archivo no puede superar {CATALOG_MAX_ROWS} productos."],
            }
        ]
    if not nonempty_rows:
        return [], [{"row": 2, "errors": ["El archivo no contiene productos."]}]

    parsed_rows = []
    errors = []
    seen_barcodes = {}
    seen_identities = {}
    seen_references = {}
    for row_number, canonical in nonempty_rows:
        parsed, error = _parse_catalog_row(canonical, row_number)
        if error:
            errors.append(error)
            continue

        if parsed["barcode"]:
            previous_row = seen_barcodes.get(parsed["barcode"])
            if previous_row:
                errors.append(
                    {
                        "row": row_number,
                        "errors": [
                            f"Código de barra repetido; también aparece en la fila {previous_row}."
                        ],
                    }
                )
                continue
            seen_barcodes[parsed["barcode"]] = row_number

        if parsed["internal_reference"]:
            reference_key = str(parsed["internal_reference"])
            previous_row = seen_references.get(reference_key)
            if previous_row:
                errors.append(
                    {
                        "row": row_number,
                        "errors": [
                            f"Referencia interna repetida; también aparece en la fila {previous_row}."
                        ],
                    }
                )
                continue
            seen_references[reference_key] = row_number

        identity = tuple(
            value.casefold()
            for value in (
                parsed["name"],
                parsed["brand"],
                parsed["presentation"],
            )
        )
        if not parsed["barcode"] and not parsed["internal_reference"]:
            previous_row = seen_identities.get(identity)
            if previous_row:
                errors.append(
                    {
                        "row": row_number,
                        "errors": [
                            f"Producto sin código repetido; también aparece en la fila {previous_row}."
                        ],
                    }
                )
                continue
            seen_identities[identity] = row_number
        parsed_rows.append(parsed)
    return parsed_rows, errors


def _catalog_product_values(row):
    return {
        "barcode": row["barcode"],
        "name": row["name"],
        "brand": row["brand"],
        "presentation": row["presentation"],
        "category": row["category"],
        "unit": row["unit"],
        "sale_mode": row["sale_mode"],
        "units_per_case": row["units_per_case"],
        "is_active": row["is_active"],
    }


def _find_catalog_product(row):
    if row["internal_reference"]:
        return MasterProduct.objects.select_for_update().filter(
            internal_reference=row["internal_reference"]
        ).first()
    if row["barcode"]:
        return MasterProduct.objects.select_for_update().filter(
            barcode=row["barcode"]
        ).first()
    matches = MasterProduct.objects.select_for_update().filter(
        name__iexact=row["name"],
        brand__iexact=row["brand"],
        presentation__iexact=row["presentation"],
    )
    if matches.count() > 1:
        raise CatalogImportAbort(
            [
                {
                    "row": row["row_number"],
                    "errors": [
                        "Hay más de un producto existente con el mismo nombre, marca y presentación."
                    ],
                }
            ]
        )
    return matches.first()


def _apply_catalog_rows(rows):
    created = 0
    updated = 0
    unchanged = 0
    for row in rows:
        product = _find_catalog_product(row)
        values = _catalog_product_values(row)
        if product is None:
            create_values = dict(values)
            if row["internal_reference"]:
                create_values["internal_reference"] = row["internal_reference"]
            MasterProduct.objects.create(**create_values)
            created += 1
            continue

        if row["barcode"]:
            barcode_owner = MasterProduct.objects.exclude(pk=product.pk).filter(
                barcode=row["barcode"]
            ).first()
            if barcode_owner is not None:
                raise CatalogImportAbort(
                    [
                        {
                            "row": row["row_number"],
                            "errors": ["El código de barra pertenece a otro producto."],
                        }
                    ]
                )

        changed_fields = []
        for field, value in values.items():
            if getattr(product, field) != value:
                setattr(product, field, value)
                changed_fields.append(field)
        if changed_fields:
            product.save(update_fields=tuple(changed_fields) + ("updated_at",))
            updated += 1
        else:
            unchanged += 1
    return created, updated, unchanged


def _finish_catalog_import(
    catalog_import,
    *,
    status_value,
    total_rows,
    created=0,
    updated=0,
    unchanged=0,
    errors=None,
):
    errors = errors or []
    catalog_import.status = status_value
    catalog_import.total_rows = total_rows
    catalog_import.created_products = created
    catalog_import.updated_products = updated
    catalog_import.unchanged_products = unchanged
    catalog_import.error_rows = len(errors)
    catalog_import.errors = errors[:CATALOG_MAX_REPORTED_ERRORS]
    catalog_import.completed_at = timezone.now()
    catalog_import.save(
        update_fields=(
            "status",
            "total_rows",
            "created_products",
            "updated_products",
            "unchanged_products",
            "error_rows",
            "errors",
            "completed_at",
            "updated_at",
        )
    )
    return catalog_import


def import_master_catalog(*, uploaded_by, uploaded_file):
    """Validate every row, then atomically create or update master products."""

    filename = Path(uploaded_file.name).name
    file_format = (
        CatalogImport.XLSX if Path(filename).suffix.lower() == ".xlsx"
        else CatalogImport.CSV
    )
    catalog_import = CatalogImport.objects.create(
        uploaded_by=uploaded_by,
        source_file=uploaded_file,
        original_filename=filename,
        file_format=file_format,
    )

    try:
        catalog_import.source_file.open("rb")
        file_bytes = catalog_import.source_file.read()
        catalog_import.source_file.close()
        if file_format == CatalogImport.XLSX:
            raw_rows, fieldnames = _read_xlsx_catalog(file_bytes)
        else:
            raw_rows, fieldnames = _read_csv_catalog(file_bytes)
        rows, errors = _validate_catalog_rows(raw_rows, fieldnames)
        if errors:
            return _finish_catalog_import(
                catalog_import,
                status_value=CatalogImport.FAILED,
                total_rows=len(rows) + len(errors),
                errors=errors,
            )

        try:
            with transaction.atomic():
                created, updated, unchanged = _apply_catalog_rows(rows)
        except CatalogImportAbort as exc:
            return _finish_catalog_import(
                catalog_import,
                status_value=CatalogImport.FAILED,
                total_rows=len(rows),
                errors=exc.errors,
            )
        except IntegrityError:
            return _finish_catalog_import(
                catalog_import,
                status_value=CatalogImport.FAILED,
                total_rows=len(rows),
                errors=[
                    {
                        "row": 1,
                        "errors": [
                            "Otro proceso modificó el catálogo. Intente importar nuevamente."
                        ],
                    }
                ],
            )

        return _finish_catalog_import(
            catalog_import,
            status_value=CatalogImport.COMPLETED,
            total_rows=len(rows),
            created=created,
            updated=updated,
            unchanged=unchanged,
        )
    except CatalogImportAbort as exc:
        return _finish_catalog_import(
            catalog_import,
            status_value=CatalogImport.FAILED,
            total_rows=0,
            errors=exc.errors,
        )


def _dashboard_sales(*, organisation, selected_date, store=None):
    money_field = DecimalField(max_digits=18, decimal_places=2)
    zero = Decimal("0.00")
    sales = Sale.objects.filter(
        organisation=organisation,
        status=Sale.COMPLETED,
        created_at__date=selected_date,
    )
    sale_items = SaleItem.objects.filter(
        sale__organisation=organisation,
        sale__status=Sale.COMPLETED,
        sale__created_at__date=selected_date,
    )
    if store is not None:
        sales = sales.filter(store=store)
        sale_items = sale_items.filter(sale__store=store)

    totals = sales.aggregate(
        sale_count=Count("id"),
        revenue=Coalesce(Sum("total"), zero, output_field=money_field),
        discounts=Coalesce(Sum("discount"), zero, output_field=money_field),
    )
    cost_of_goods = sale_items.aggregate(
        value=Coalesce(
            Sum(F("quantity") * F("unit_cost")),
            zero,
            output_field=money_field,
        )
    )["value"]
    payment_rows = sales.values("payment_method").annotate(
        amount=Coalesce(Sum("total"), zero, output_field=money_field),
        count=Count("id"),
    )
    payments = {
        payment_method: {"amount": "0.00", "count": 0}
        for payment_method, _label in Sale.PAYMENT_METHOD_CHOICES
    }
    for row in payment_rows:
        payments[row["payment_method"]] = {
            "amount": str(_money(row["amount"])),
            "count": row["count"],
        }

    return {
        "sale_count": totals["sale_count"],
        "revenue": _money(totals["revenue"]),
        "discounts": _money(totals["discounts"]),
        "cost_of_goods": _money(cost_of_goods),
        "payments": payments,
        "sale_items": sale_items,
    }


def build_dashboard(*, organisation, selected_date, store=None):
    """Build one compact, Spanish-ready daily summary for an owner."""

    today_sales = _dashboard_sales(
        organisation=organisation,
        selected_date=selected_date,
        store=store,
    )
    previous_sales = _dashboard_sales(
        organisation=organisation,
        selected_date=selected_date - timedelta(days=1),
        store=store,
    )
    revenue = today_sales["revenue"]
    cost_of_goods = today_sales["cost_of_goods"]
    gross_profit = _money(revenue - cost_of_goods)

    expenses = Expense.objects.filter(
        organisation=organisation,
        expense_date=selected_date,
    )
    payroll = PayrollPayment.objects.filter(
        organisation=organisation,
        paid_on=selected_date,
    )
    if store is not None:
        expenses = expenses.filter(store=store)
        payroll = payroll.filter(store=store)
    operating_expenses = _money(
        expenses.aggregate(value=Sum("amount"))["value"] or Decimal("0.00")
    )
    payroll_total = _money(
        payroll.aggregate(value=Sum("amount"))["value"] or Decimal("0.00")
    )
    net_profit = _money(gross_profit - operating_expenses - payroll_total)

    previous_revenue = previous_sales["revenue"]
    change_amount = _money(revenue - previous_revenue)
    change_percent = None
    if previous_revenue > 0:
        change_percent = _money((change_amount / previous_revenue) * 100)

    inventory = InventoryItem.objects.select_related(
        "store",
        "master_product",
    ).filter(
        organisation=organisation,
        is_active=True,
        store__is_active=True,
        master_product__is_active=True,
    )
    if store is not None:
        inventory = inventory.filter(store=store)
    low_stock = inventory.filter(quantity__lte=F("reorder_level")).order_by(
        "quantity",
        "master_product__name",
    )
    low_stock_items = [
        {
            "inventory_item_id": item.id,
            "store_id": item.store_id,
            "store_name": item.store.name,
            "product_name": str(item.master_product),
            "barcode": item.master_product.barcode,
            "quantity": str(item.quantity),
            "reorder_level": str(item.reorder_level),
        }
        for item in low_stock[:5]
    ]

    active_orders = CustomerOrder.objects.filter(
        organisation=organisation,
        status__in=(
            CustomerOrder.NEW,
            CustomerOrder.PREPARING,
            CustomerOrder.READY,
            CustomerOrder.ON_THE_WAY,
        ),
    )
    if store is not None:
        active_orders = active_orders.filter(store=store)
    order_counts = {
        row["status"]: row["count"]
        for row in active_orders.values("status").annotate(count=Count("id"))
    }
    recent_orders = [
        {
            "id": order.id,
            "order_number": str(order.order_number),
            "store_id": order.store_id,
            "store_name": order.store.name,
            "customer_name": order.customer_name,
            "status": order.status,
            "status_display": order.get_status_display(),
            "total": str(order.total),
            "created_at": order.created_at,
        }
        for order in active_orders.select_related("store").order_by("created_at")[:5]
    ]

    top_products = list(
        today_sales["sale_items"]
        .values("product_name")
        .annotate(
            quantity_sold=Sum("quantity"),
            sales_total=Sum("line_total"),
        )
        .order_by("-quantity_sold", "product_name")[:5]
    )
    for product in top_products:
        product["quantity_sold"] = str(
            Decimal(product["quantity_sold"]).quantize(Decimal("0.001"))
        )
        product["sales_total"] = str(_money(product["sales_total"]))

    outstanding_credit = _money(
        Customer.objects.filter(
            organisation=organisation,
            is_active=True,
        ).aggregate(value=Sum("balance"))["value"]
        or Decimal("0.00")
    )

    store_comparison = []
    if store is None:
        for current_store in Store.objects.filter(
            organisation=organisation,
            is_active=True,
        ).order_by("name"):
            store_sales = _dashboard_sales(
                organisation=organisation,
                selected_date=selected_date,
                store=current_store,
            )
            store_comparison.append(
                {
                    "store_id": current_store.id,
                    "store_name": current_store.name,
                    "sale_count": store_sales["sale_count"],
                    "revenue": str(store_sales["revenue"]),
                    "gross_profit": str(
                        _money(
                            store_sales["revenue"]
                            - store_sales["cost_of_goods"]
                        )
                    ),
                }
            )

    return {
        "date": str(selected_date),
        "store": (
            {"id": store.id, "name": store.name}
            if store is not None
            else None
        ),
        "summary": {
            "sale_count": today_sales["sale_count"],
            "revenue": str(revenue),
            "discounts": str(today_sales["discounts"]),
            "cost_of_goods": str(cost_of_goods),
            "gross_profit": str(gross_profit),
            "operating_expenses": str(operating_expenses),
            "payroll": str(payroll_total),
            "net_profit": str(net_profit),
            "is_profitable": net_profit >= 0,
            "outstanding_credit": str(outstanding_credit),
            "low_stock_count": low_stock.count(),
            "new_orders": order_counts.get(CustomerOrder.NEW, 0),
            "active_orders": sum(order_counts.values()),
        },
        "comparison": {
            "previous_date": str(selected_date - timedelta(days=1)),
            "previous_revenue": str(previous_revenue),
            "change_amount": str(change_amount),
            "change_percent": (
                str(change_percent) if change_percent is not None else None
            ),
        },
        "payments": today_sales["payments"],
        "low_stock": low_stock_items,
        "orders": recent_orders,
        "top_products": top_products,
        "stores": store_comparison,
    }
